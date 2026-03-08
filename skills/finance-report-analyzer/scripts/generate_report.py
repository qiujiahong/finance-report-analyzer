#!/usr/bin/env python3
"""
Generate interactive HTML financial analysis report from parsed Excel data.

Usage:
    python3 generate_report.py <input_xlsx> <output_html> [--company NAME] [--ticker TICKER]

The script:
1. Reads Excel file with financial summary data
2. Extracts annual report columns + forecast columns
3. Generates a self-contained HTML report with analysis
"""

import sys
import os
import argparse
import json

try:
    import openpyxl
except ImportError:
    os.system("pip install openpyxl -q")
    import openpyxl


def parse_excel(filepath):
    """Parse financial Excel file into structured data."""
    wb = openpyxl.load_workbook(filepath, data_only=True)
    ws = wb[wb.sheetnames[0]]

    rows = {}
    header_row = None
    for row in ws.iter_rows(min_row=1, values_only=False):
        vals = [cell.value for cell in row]
        label = str(vals[0]).strip() if vals[0] else ""
        if not header_row:
            header_row = vals[1:]
        rows[label] = vals[1:]

    return rows, header_row


def identify_periods(rows):
    """Identify report period types and select annual + forecast columns."""
    report_types = []
    for key in rows:
        if "报告期" in key or "报表" in key:
            vals = [str(v) if v else "" for v in rows[key]]
            if any("年报" in v or "季报" in v or "中报" in v for v in vals):
                report_types = vals
                break

    if not report_types:
        # Fallback: use all columns
        first_key = list(rows.keys())[0]
        return list(range(len(rows[first_key]))), ["" for _ in rows[first_key]]

    annual_idx = [i for i, r in enumerate(report_types) if r.endswith("年报") and "下半年" not in r]
    forecast_idx = [i for i, r in enumerate(report_types) if "盈利预测" in r]
    all_idx = annual_idx + forecast_idx

    return all_idx, report_types


def get_values(rows, label, indices):
    """Extract values for specific column indices."""
    vals = rows.get(label, [])
    return [vals[i] if i < len(vals) else None for i in indices]


def fmt(v, decimals=2):
    if v is None or v == "" or str(v) == "None":
        return "-"
    try:
        f = float(v)
        if abs(f) > 1000:
            return f"{f:,.{decimals}f}"
        return f"{f:.{decimals}f}"
    except (ValueError, TypeError):
        return str(v)


def fmt_pct(v):
    if v is None or v == "" or str(v) == "None":
        return "-"
    try:
        return f"{float(v):.1f}%"
    except (ValueError, TypeError):
        return str(v)


def find_row_label(rows, keywords):
    """Find a row label containing any of the keywords."""
    for label in rows:
        for kw in keywords:
            if kw in label:
                return label
    return None


def generate_html(rows, header_row, indices, report_types, company="", ticker=""):
    """Generate the full HTML report."""

    # Extract dates from header
    dates = []
    for i in indices:
        h = str(header_row[i]) if i < len(header_row) and header_row[i] else ""
        dates.append(h[:4] if len(h) >= 4 else h)

    is_forecast = ["盈利预测" in report_types[i] if i < len(report_types) else False for i in indices]

    # Metric mapping - try multiple possible labels
    def get(keywords):
        label = find_row_label(rows, keywords)
        if label:
            return get_values(rows, label, indices)
        return [None] * len(indices)

    revenue = get(["营业总收入", "营业收入"])
    op_profit = get(["营业利润"])
    net_profit_parent = get(["归属母公司股东的净利润", "归母净利润"])
    net_profit = get(["净利润"])
    rd = get(["研发支出", "研发费用"])
    ebitda = get(["EBITDA"])
    gross_margin = get(["销售毛利率", "毛利率"])
    net_margin = get(["销售净利率", "净利率"])
    roe = get(["ROE(摊薄)", "ROE"])
    roa = get(["ROA"])
    total_assets = get(["资产总计", "总资产"])
    total_liab = get(["负债合计", "总负债"])
    equity = get(["归属母公司股东的权益", "股东权益"])
    current_assets = get(["流动资产"])
    debt_ratio = get(["资产负债率"])
    ocf = get(["经营活动现金净流量", "经营活动现金流"])
    icf = get(["投资活动现金净流量"])
    fcf_finance = get(["筹资活动现金净流量"])
    cash_end = get(["期末现金余额"])
    capex = get(["购建固定无形长期资产支付的现金", "资本开支"])
    eps = get(["EPS(基本)", "EPS"])
    bps = get(["每股净资产", "BPS"])
    employees = get(["员工总数"])
    asset_turnover = get(["资产周转率"])

    # Compute free cash flow
    fcf = []
    for o, c in zip(ocf, capex):
        if o is not None and c is not None:
            try:
                fcf.append(float(o) - float(c))
            except (ValueError, TypeError):
                fcf.append(None)
        else:
            fcf.append(None)

    title = f"{company} [{ticker}]" if ticker else company or "财务分析报告"

    def table_header():
        h = "<tr><th>指标</th>"
        for i, y in enumerate(dates):
            cls = ' class="forecast"' if is_forecast[i] else ""
            suffix = "E" if is_forecast[i] else ""
            h += f"<th{cls}>{y}{suffix}</th>"
        return h + "</tr>"

    def table_row(label, values, is_pct=False):
        r = f"<tr><td>{label}</td>"
        for i, v in enumerate(values):
            cls = ' class="forecast"' if is_forecast[i] else ""
            if is_pct:
                r += f"<td{cls}>{fmt_pct(v)}</td>"
            else:
                try:
                    fv = float(v) if v is not None else None
                    color = ""
                    if fv is not None and fv > 0:
                        color = ' class="positive"'
                    elif fv is not None and fv < 0:
                        color = ' class="negative"'
                    r += f"<td{cls}><span{color}>{fmt(v)}</span></td>"
                except (ValueError, TypeError):
                    r += f"<td{cls}>{fmt(v)}</td>"
        return r + "</tr>"

    # Latest actual year values for summary cards
    actual_idx = [i for i, f in enumerate(is_forecast) if not f]
    latest = actual_idx[-1] if actual_idx else 0
    forecast_first = next((i for i, f in enumerate(is_forecast) if f), None)

    html = f"""<!DOCTYPE html>
<html lang="zh-CN">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>{title} 财务分析报告</title>
<style>
:root {{ --primary: #1a56db; --danger: #dc2626; --success: #16a34a; --bg: #f8fafc; --card: #fff; --border: #e2e8f0; --text: #1e293b; --text2: #64748b; }}
* {{ margin:0; padding:0; box-sizing:border-box; }}
body {{ font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, sans-serif; background:var(--bg); color:var(--text); line-height:1.6; }}
.container {{ max-width:1200px; margin:0 auto; padding:20px; }}
.header {{ background:linear-gradient(135deg,#1e3a8a,#3b82f6); color:#fff; padding:40px; border-radius:16px; margin-bottom:24px; }}
.header h1 {{ font-size:28px; margin-bottom:8px; }}
.header .subtitle {{ opacity:0.85; font-size:14px; }}
.card {{ background:var(--card); border-radius:12px; padding:24px; margin-bottom:20px; box-shadow:0 1px 3px rgba(0,0,0,0.08); }}
.card h2 {{ font-size:18px; margin-bottom:16px; padding-bottom:8px; border-bottom:2px solid var(--primary); display:inline-block; }}
.metrics-grid {{ display:grid; grid-template-columns:repeat(auto-fit,minmax(200px,1fr)); gap:16px; margin-bottom:20px; }}
.metric-card {{ background:#f1f5f9; border-radius:10px; padding:16px; text-align:center; }}
.metric-card .label {{ font-size:12px; color:var(--text2); margin-bottom:4px; }}
.metric-card .value {{ font-size:24px; font-weight:700; }}
.metric-card .change {{ font-size:12px; margin-top:4px; }}
.positive {{ color:var(--success); }}
.negative {{ color:var(--danger); }}
table {{ width:100%; border-collapse:collapse; font-size:13px; overflow-x:auto; display:block; }}
th, td {{ padding:8px 10px; text-align:right; white-space:nowrap; border-bottom:1px solid var(--border); }}
th {{ background:#f1f5f9; font-weight:600; position:sticky; top:0; }}
td:first-child, th:first-child {{ text-align:left; font-weight:500; position:sticky; left:0; background:#fff; z-index:1; min-width:180px; }}
th:first-child {{ background:#f1f5f9; z-index:2; }}
tr:hover td {{ background:#f8fafc; }}
.forecast {{ background:#fffbeb !important; }}
.summary-box {{ background:linear-gradient(135deg,#eff6ff,#dbeafe); border-left:4px solid var(--primary); padding:16px 20px; border-radius:0 8px 8px 0; margin:16px 0; }}
.two-col {{ display:grid; grid-template-columns:1fr 1fr; gap:20px; }}
@media (max-width:768px) {{ .two-col {{ grid-template-columns:1fr; }} }}
.tag-good {{ display:inline-block; padding:2px 8px; border-radius:4px; font-size:11px; background:#dcfce7; color:#166534; }}
</style>
</head>
<body>
<div class="container">
<div class="header">
<h1>📊 {title}</h1>
<div class="subtitle">财务分析报告 | 单位：亿元(CNY)</div>
</div>

<div class="card">
<h2>📋 核心摘要</h2>
<div class="metrics-grid">
<div class="metric-card">
<div class="label">{dates[latest]}年营收</div>
<div class="value">{fmt(revenue[latest])}</div>
</div>
<div class="metric-card">
<div class="label">{dates[latest]}年归母净利润</div>
<div class="value">{fmt(net_profit_parent[latest])}</div>
</div>
<div class="metric-card">
<div class="label">{dates[latest]}年毛利率</div>
<div class="value">{fmt_pct(gross_margin[latest])}</div>
</div>
{"<div class='metric-card'><div class='label'>" + dates[forecast_first] + "E归母净利润</div><div class='value'>" + fmt(net_profit_parent[forecast_first]) + "</div></div>" if forecast_first is not None else ""}
</div>
</div>

<div class="card">
<h2>📈 一、盈利能力分析</h2>
<table>
{table_header()}
{table_row("营业总收入", revenue)}
{table_row("营业利润", op_profit)}
{table_row("归母净利润", net_profit_parent)}
{table_row("研发支出", rd)}
{table_row("EBITDA", ebitda)}
</table>
<h3 style="margin:16px 0 8px;font-size:15px;">利润率指标</h3>
<table>
{table_header()}
{table_row("毛利率", gross_margin, True)}
{table_row("净利率", net_margin, True)}
{table_row("ROE(摊薄)", roe, True)}
{table_row("ROA", roa, True)}
</table>
</div>

<div class="card">
<h2>🏦 二、资产负债分析</h2>
<table>
{table_header()}
{table_row("总资产", total_assets)}
{table_row("流动资产", current_assets)}
{table_row("总负债", total_liab)}
{table_row("股东权益", equity)}
{table_row("资产负债率", debt_ratio, True)}
</table>
</div>

<div class="card">
<h2>💰 三、现金流分析</h2>
<table>
{table_header()}
{table_row("经营活动现金流", ocf)}
{table_row("投资活动现金流", icf)}
{table_row("筹资活动现金流", fcf_finance)}
{table_row("资本开支", capex)}
{table_row("自由现金流(OCF-CapEx)", fcf)}
{table_row("期末现金余额", cash_end)}
</table>
</div>

<div class="card">
<h2>📊 四、每股指标与效率</h2>
<table>
{table_header()}
{table_row("EPS(元)", eps)}
{table_row("每股净资产(元)", bps)}
{table_row("资产周转率(倍)", asset_turnover)}
{table_row("员工人数", employees)}
</table>
</div>

<div style="text-align:center;color:var(--text2);font-size:12px;padding:20px;">
📅 数据来源：公司财务报告 | 仅供参考，不构成投资建议
</div>
</div>
</body>
</html>"""

    return html


def main():
    parser = argparse.ArgumentParser(description="Generate financial analysis HTML report")
    parser.add_argument("input", help="Input Excel file path")
    parser.add_argument("output", help="Output HTML file path")
    parser.add_argument("--company", default="", help="Company name")
    parser.add_argument("--ticker", default="", help="Stock ticker")
    args = parser.parse_args()

    rows, header_row = parse_excel(args.input)
    indices, report_types = identify_periods(rows)
    html = generate_html(rows, header_row, indices, report_types, args.company, args.ticker)

    with open(args.output, "w", encoding="utf-8") as f:
        f.write(html)

    print(f"Report generated: {args.output} ({len(html)} bytes)")


if __name__ == "__main__":
    main()
